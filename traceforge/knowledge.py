import asyncio
import hashlib
import io
import json
import time
from pathlib import Path
from bs4 import BeautifulSoup
import httpx
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter
from .config import settings
from .db import db,uid
from .providers import embeddings, chat_model
from .security import public_url,redact

def split_document(text, metadata):
    sections=MarkdownHeaderTextSplitter(headers_to_split_on=[('#','h1'),('##','h2'),('###','h3')],strip_headers=False).split_text(text)
    for d in sections:d.metadata=metadata|d.metadata
    split=RecursiveCharacterTextSplitter(chunk_size=1000,chunk_overlap=150,separators=['\n\n','\n','。','；',' ',''],add_start_index=True)
    docs=split.split_documents(sections or [Document(page_content=text,metadata=metadata)])
    for d in docs:
        offset=text.find(d.page_content)
        d.metadata['line']=text[:max(0,offset)].count('\n')+1
    return docs

def parse_file(filename,content):
    suffix=Path(filename).suffix.lower()
    if suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(content))
        pages=[(p.extract_text() or '') for p in reader.pages]
        if not ''.join(pages).strip():raise ValueError('扫描 PDF 未发现可提取文字，请上传 OCR 后文本或截图')
        return '\n\n'.join(f'## 第 {i+1} 页\n{p}' for i,p in enumerate(pages))
    if suffix=='.docx':
        from docx import Document as WordDocument
        doc=WordDocument(io.BytesIO(content))
        return '\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    if suffix not in {'.md','.txt','.log','.json','.yaml','.yml','.csv','.py','.js','.ts','.xml','.conf'}:
        raise ValueError('支持 Markdown / 文本 / 日志 / JSON / YAML / CSV / PDF / DOCX / 图片')
    try:return content.decode('utf-8-sig')
    except UnicodeDecodeError:return content.decode('gb18030')

async def fetch_public(url):
    public_url(url)
    async with httpx.AsyncClient(trust_env=False,timeout=25,follow_redirects=False,headers={'User-Agent':'TraceForge/0.1 knowledge import'}) as c:
        for _ in range(5):
            public_url(url)
            r=await c.get(url)
            if r.is_redirect:
                url=str(r.url.join(r.headers['location']));continue
            r.raise_for_status()
            if len(r.content)>5_000_000:raise ValueError('网页超过 5MB 限制')
            if r.url.host=='api.github.com' and '/readme' in r.url.path:
                import base64
                payload=r.json()
                return payload.get('name','GitHub README'),base64.b64decode(payload['content']).decode('utf-8'),payload.get('html_url',url)
            if 'html' in r.headers.get('content-type',''):
                soup=BeautifulSoup(r.text,'html.parser')
                for x in soup(['script','style','nav','footer','header']):x.decompose()
                title=soup.title.get_text(strip=True) if soup.title else url
                body=(soup.find('main') or soup.find('article') or soup).get_text('\n',strip=True)
                return title,body,url
            return Path(url.split('?')[0]).name,r.text,url
    raise ValueError('网页重定向次数过多')

class Knowledge:
    def __init__(self):
        self.client=None
        self.lock=asyncio.Lock()

    def connect(self):
        from pymilvus import MilvusClient,DataType,Function,FunctionType
        if self.client:return self.client
        c=MilvusClient(uri=settings.milvus_uri,timeout=20)
        if not c.has_collection(settings.collection):
            schema=c.create_schema(auto_id=False,enable_dynamic_field=False)
            schema.add_field('pk',DataType.VARCHAR,is_primary=True,max_length=64)
            schema.add_field('text',DataType.VARCHAR,max_length=16000,enable_analyzer=True,analyzer_params={'type':'chinese'})
            for name,size in [('document_id',64),('workspace_id',64),('title',1024),('source',4096),('stack',64),('version',128)]:schema.add_field(name,DataType.VARCHAR,max_length=size)
            schema.add_field('metadata',DataType.JSON)
            schema.add_field('dense',DataType.FLOAT_VECTOR,dim=1024)
            schema.add_field('sparse',DataType.SPARSE_FLOAT_VECTOR)
            schema.add_function(Function(name='bm25_text',input_field_names=['text'],output_field_names=['sparse'],function_type=FunctionType.BM25))
            indexes=c.prepare_index_params()
            indexes.add_index('dense',index_type='HNSW',metric_type='COSINE',params={'M':16,'efConstruction':100})
            indexes.add_index('sparse',index_type='SPARSE_INVERTED_INDEX',metric_type='BM25')
            for field in ('workspace_id','stack','document_id'):indexes.add_index(field,index_type='INVERTED')
            c.create_collection(settings.collection,schema=schema,index_params=indexes,consistency_level='Strong')
        self.client=c
        return c

    async def ingest(self,record):
        async with self.lock:
            record['status']='indexing';db.put('document',record['id'],record)
            try:
                docs=split_document(record['text'],{'document_id':record['id'],'source':record['source']})
                vectors=await embeddings().aembed_documents([d.page_content for d in docs])
                if any(len(v)!=1024 for v in vectors):raise ValueError('嵌入维度不匹配，必须为1024')
                rows=[]
                for i,(doc,vector) in enumerate(zip(docs,vectors)):
                    rows.append({'pk':hashlib.sha256(f'{record["id"]}:{i}'.encode()).hexdigest(),'text':doc.page_content,'dense':vector,'document_id':record['id'],'workspace_id':record.get('workspace_id','global'),'stack':record.get('stack','general'),'version':record.get('version',''),'title':record['title'][:300],'source':record['source'][:3000],'metadata':doc.metadata})
                client=await asyncio.to_thread(self.connect)
                await asyncio.to_thread(client.delete,collection_name=settings.collection,filter='document_id == {doc}',filter_params={'doc':record['id']})
                for start in range(0,len(rows),64):await asyncio.to_thread(client.upsert,collection_name=settings.collection,data=rows[start:start+64])
                record.update(status='ready',chunks=len(rows),indexed_at=time.time(),error=None)
            except Exception as e:record.update(status='failed',error=redact(str(e))[:800])
            db.put('document',record['id'],record)
            return record

    def add(self,title,text,source,stack='general',workspace_id='global',version=''):
        text=redact(text)
        id=hashlib.sha256((source+'\0'+workspace_id).encode()).hexdigest()[:32]
        previous=db.get('document',id)
        checksum=hashlib.sha256(text.encode()).hexdigest()
        if previous and previous.get('checksum')==checksum and previous['status']=='ready':return previous
        record={'id':id,'title':title,'source':source,'text':text,'stack':stack,'workspace_id':workspace_id,'version':version,'checksum':checksum,'status':'queued','chunks':0,'created':time.time()}
        db.put('document',id,record)
        return record

    async def delete(self,id):
        client=await asyncio.to_thread(self.connect)
        await asyncio.to_thread(client.delete,collection_name=settings.collection,filter='document_id == {doc}',filter_params={'doc':id})
        db.delete('document',id)

    async def screenshot(self,filename,content):
        import base64
        from langchain_core.messages import HumanMessage
        mime={'.png':'image/png','.jpg':'image/jpeg','.jpeg':'image/jpeg','.webp':'image/webp'}[Path(filename).suffix.lower()]
        response=await chat_model().ainvoke([HumanMessage(content=[{'type':'text','text':'提取截图中的完整报错、错误码、文件名和行号。仅记录可见内容，不猜测。'}, {'type':'image_url','image_url':{'url':f'data:{mime};base64,'+base64.b64encode(content).decode()}}])])
        return redact(str(response.content))

knowledge=Knowledge()
